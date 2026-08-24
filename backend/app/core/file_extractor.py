import os
import io
import json
import logging
from pathlib import Path
from typing import Dict, Any, Optional

logger = logging.getLogger("workbench.extractor")

def extract_file_content(file_path: str, max_preview_len: int = 2000) -> Dict[str, Any]:
    """
    Extract text, tables count, and preview from an uploaded document.
    Supports .pdf, .docx, .xlsx, .csv, .txt, .log, .json, .md, and image metadata.
    """
    path = Path(file_path)
    if not path.exists():
        logger.warning(f"File not found for extraction: {file_path}")
        return {
            "filename": path.name,
            "text": "",
            "preview": "",
            "tables_found": 0,
            "text_length": 0,
            "file_type": "unknown",
            "error": "File does not exist on disk"
        }

    ext = path.suffix.lower()
    filename = path.name
    extracted_text = ""
    tables_found = 0
    file_type = ext.replace(".", "")

    try:
        if ext in (".txt", ".md", ".log", ".json", ".yaml", ".yml"):
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                extracted_text = f.read()

        elif ext == ".csv":
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                lines = f.readlines()
                tables_found = 1
                extracted_text = "".join(lines[:100])

        elif ext in (".xlsx", ".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                tables_found = len(wb.sheetnames)
                sheet_summaries = []
                for sheet_name in wb.sheetnames:
                    sheet = wb[sheet_name]
                    rows_sample = []
                    for i, row in enumerate(sheet.iter_rows(values_only=True)):
                        if i >= 30:  # limit rows sample
                            break
                        row_vals = [str(val) if val is not None else "" for val in row]
                        if any(row_vals):
                            rows_sample.append(" | ".join(row_vals))
                    sheet_summaries.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows_sample))
                extracted_text = "\n\n".join(sheet_summaries)
            except Exception as e:
                logger.warning(f"openpyxl failed to read {path}: {e}")
                extracted_text = f"[Excel File: {filename} - {os.path.getsize(path)} bytes]"

        elif ext == ".docx":
            try:
                import docx
                doc = docx.Document(path)
                tables_found = len(doc.tables)
                paras = [p.text for p in doc.paragraphs if p.text.strip()]
                extracted_text = "\n".join(paras)
                if doc.tables:
                    table_texts = []
                    for t_idx, table in enumerate(doc.tables):
                        t_rows = []
                        for row in table.rows:
                            row_cells = [cell.text.strip() for cell in row.cells]
                            t_rows.append(" | ".join(row_cells))
                        table_texts.append(f"--- Table {t_idx+1} ---\n" + "\n".join(t_rows[:20]))
                    extracted_text += "\n\n" + "\n\n".join(table_texts)
            except Exception as e:
                logger.warning(f"python-docx failed to read {path}: {e}")
                extracted_text = f"[Word Document: {filename} - {os.path.getsize(path)} bytes]"

        elif ext == ".pdf":
            try:
                import pypdf
                reader = pypdf.PdfReader(str(path))
                pages_text = []
                for i, page in enumerate(reader.pages):
                    if i >= 10:  # first 10 pages
                        break
                    p_text = page.extract_text()
                    if p_text:
                        pages_text.append(f"--- Page {i+1} ---\n{p_text}")
                extracted_text = "\n\n".join(pages_text)
            except Exception as e:
                logger.warning(f"pypdf extraction failed on {path}: {e}")
                extracted_text = f"[PDF Document: {filename} - {os.path.getsize(path)} bytes]"

        elif ext in (".png", ".jpg", ".jpeg", ".bmp", ".tiff"):
            # Image file metadata
            file_type = "image"
            size_kb = round(os.path.getsize(path) / 1024, 1)
            extracted_text = f"[Image Scan: {filename}, size: {size_kb} KB. Visual inspection scan attached.]"

        else:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                extracted_text = f.read(5000)

    except Exception as e:
        logger.error(f"Error extracting content from {file_path}: {e}")
        extracted_text = f"[File Attachment: {filename} ({os.path.getsize(path)} bytes)]"

    preview = extracted_text[:max_preview_len].strip()
    if len(extracted_text) > max_preview_len:
        preview += "..."

    return {
        "filename": filename,
        "text": extracted_text,
        "preview": preview,
        "tables_found": tables_found,
        "text_length": len(extracted_text),
        "file_type": file_type
    }
